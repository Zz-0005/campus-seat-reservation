#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 docs 下的 PRD Markdown 渲染成排版好的 Word 文档。

用法：
    python tools/build_prd_docx.py

只支持本项目 PRD 用到的 Markdown 子集：标题、段落、无序/有序列表、
表格、代码块、引用、加粗与行内代码。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "PRD-校园自习座位预约系统.md"
TARGET = ROOT / "docs" / "PRD-校园自习座位预约系统.docx"

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x38, 0x64)
ACCENT_SOFT = RGBColor(0x2F, 0x54, 0x96)
HEADER_FILL = "E8EEF9"


# ---------------------------------------------------------------- 工具函数
def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None, mono: bool = False):
    font_name = MONO_FONT if mono else name
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_of(doc, name: str):
    return doc.styles[name]


def set_style_eastasia(style, font: str):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)


def configure_styles(doc: Document):
    normal = style_of(doc, "Normal")
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    set_style_eastasia(normal, BODY_FONT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)

    specs = {
        "Title": (HEAD_FONT, 24, True, ACCENT, 0, 10),
        "Heading 1": (HEAD_FONT, 16, True, ACCENT, 18, 8),
        "Heading 2": (HEAD_FONT, 13, True, ACCENT_SOFT, 12, 6),
        "Heading 3": (HEAD_FONT, 11.5, True, ACCENT_SOFT, 10, 4),
        "List Bullet": (BODY_FONT, 10.5, False, None, 0, 2),
        "Quote": (BODY_FONT, 10.5, False, None, 0, 4),
    }
    for name, (font, size, bold, color, before, after) in specs.items():
        style = style_of(doc, name)
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color
        set_style_eastasia(style, font)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def shade_cell(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def mark_header_row(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def add_page_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x66, 0x66, 0x66))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    holder = paragraph.add_run()
    holder._r.append(fld_begin)
    holder._r.append(instr)
    holder._r.append(fld_end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, BODY_FONT, 9, color=RGBColor(0x66, 0x66, 0x66))


def add_header_text(section, text: str):
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    set_run_font(run, HEAD_FONT, 8.5, color=RGBColor(0x88, 0x8C, 0x99))


def add_toc(doc: Document):
    # 目录标题不使用 Heading 样式，避免它自己出现在目录里
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(8)
    run = heading.add_run("目录")
    set_run_font(run, HEAD_FONT, 16, bold=True, color=ACCENT)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "（在 Word 中右键此处选择“更新域”即可生成目录）"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)
    set_run_font(run, BODY_FONT, 10.5, color=RGBColor(0x88, 0x8C, 0x99))

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- 行内格式
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, base_size: float = 10.5, bold_all: bool = False):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            set_run_font(run, BODY_FONT, base_size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            set_run_font(run, MONO_FONT, base_size - 0.5, mono=True)
        else:
            run = paragraph.add_run(chunk)
            set_run_font(run, BODY_FONT, base_size, bold=True if bold_all else None)


# ---------------------------------------------------------------- 表格
def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]], total_width_cm: float = 16.0):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    weights = []
    for index in range(len(header)):
        longest = max(
            [len(header[index])] + [len(row[index]) if index < len(row) else 0 for row in body]
        )
        weights.append(max(longest, 4))
    total_weight = sum(weights)
    widths = [max(1.6, total_width_cm * weight / total_weight) for weight in weights]
    scale = total_width_cm / sum(widths)
    widths = [width * scale for width in widths]

    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(1)
        add_inline(paragraph, text, base_size=9.5, bold_all=True)
        shade_cell(cell, HEADER_FILL)
    mark_header_row(table.rows[0])

    for row_data in body:
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.width = Cm(widths[index])
            text = row_data[index] if index < len(row_data) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.2
            add_inline(paragraph, text, base_size=9.5)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.line_spacing = 1.0
    return table


# ---------------------------------------------------------------- 主流程
def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    add_header_text(section, "校园自习座位预约系统 产品需求文档 V1.0")
    add_page_footer(section)

    title_done = False
    toc_done = False
    index = 0
    stats = {"tables": 0, "headings": 0, "code": 0}

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # 代码块
        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            for code_line in block:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.left_indent = Cm(0.4)
                run = paragraph.add_run(code_line if code_line.strip() else " ")
                set_run_font(run, MONO_FONT, 9, mono=True)
                shade_cell_like_paragraph(paragraph)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            stats["code"] += 1
            continue

        # 表格
        if stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = split_row(lines[index])
                if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in row):
                    rows.append(row)
                index += 1
            if rows:
                add_table(doc, rows)
                stats["tables"] += 1
            continue

        # 标题
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not title_done:
                paragraph = doc.add_paragraph(text, style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_done = True
                stats["headings"] += 1
                index += 1
                continue
            if not toc_done and level == 1 and re.match(r"^\d+\.", text):
                add_toc(doc)
                toc_done = True
            style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
            doc.add_paragraph(text, style=style)
            stats["headings"] += 1
            index += 1
            continue

        # 无序列表
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        # 有序列表（保留原文编号，避免自动重新编号）
        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, "%s. %s" % (ordered.group(1), ordered.group(2)))
            index += 1
            continue

        # 引用
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph(style="Quote")
            add_inline(paragraph, stripped.lstrip("> ").strip())
            index += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            index += 1
            continue

        # 普通段落
        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    if not toc_done:
        pass

    doc.core_properties.title = "校园自习座位预约系统 产品需求文档"
    doc.core_properties.subject = "PRD V1.0"
    doc.core_properties.author = "校园自习座位预约系统项目组"
    doc.save(TARGET)

    print("已生成：%s" % TARGET)
    print("  标题 %d 个、表格 %d 张、代码块 %d 段" % (stats["headings"], stats["tables"], stats["code"]))


def shade_cell_like_paragraph(paragraph):
    """给代码块所在段落加浅灰底纹。"""
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F6FA")
    ppr.append(shd)


if __name__ == "__main__":
    build()
    sys.exit(0)
