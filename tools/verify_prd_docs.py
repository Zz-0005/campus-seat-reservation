#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验 PRD 的两份产物：Markdown 与 Word 的结构、内容一致性。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "docs" / "PRD-校园自习座位预约系统.md"
DOCX = ROOT / "docs" / "PRD-校园自习座位预约系统.docx"

KEYS = [
    "校园自习座位预约系统",
    "FR-04 预约冲突控制（并发唯一）",
    "reservation_slots",
    "SLOT_TAKEN",
    "38 项检查全部通过",
    "16 个请求同时抢同一座位同一时段",
    "146 个座位",
    "非功能需求",
    "风险与对策",
]

MOJIBAKE = "锛鏍澹璁鐢棰勭"

ok = True


def fail(msg: str):
    global ok
    ok = False
    print("  FAIL  " + msg)


def main() -> int:
    print("[Markdown]")
    md = MD.read_text(encoding="utf-8")
    print("  字符数 %d，行数 %d" % (len(md), md.count("\n") + 1))
    bad = [ch for ch in MOJIBAKE if ch in md]
    if bad:
        fail("存在乱码字符：%s" % "".join(bad))
    else:
        print("  PASS  无乱码字符")
    h1 = re.findall(r"(?m)^# (.+)$", md)
    print("  一级标题 %d 个：%s" % (len(h1), " / ".join(h1[:4]) + " …"))
    if len(h1) < 15:
        fail("一级标题少于 15 个")
    tables_md = len(re.findall(r"(?m)^\| --- ", md))
    print("  Markdown 表格 %d 张" % tables_md)
    for key in KEYS:
        if key not in md:
            fail("Markdown 缺少内容：%s" % key)

    print("\n[Word]")
    doc = Document(str(DOCX))
    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    print("  段落 %d 个，标题 %d 个，表格 %d 张" % (len(doc.paragraphs), len(headings), len(doc.tables)))

    titles = [p.text.strip() for p in headings]
    print("  标题示例：%s" % " / ".join(titles[:5]))
    if "校园自习座位预约系统 产品需求文档（PRD）" not in titles:
        fail("缺少文档标题")
    for chapter in ("1. 背景与问题", "5. 功能需求详述", "10. 验收与测试", "13. 附录"):
        if chapter not in titles:
            fail("缺少章节：%s" % chapter)

    full_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            full_text += "\n" + " | ".join(c.text for c in row.cells)

    leftovers = [token for token in ("**", "```", "| ---") if token in full_text]
    if leftovers:
        fail("Word 正文残留 Markdown 记号：%s" % ", ".join(leftovers))
    else:
        print("  PASS  正文无 Markdown 残留记号")

    for key in KEYS:
        if key not in full_text:
            fail("Word 缺少内容：%s" % key)

    empty_header = 0
    for index, table in enumerate(doc.tables, start=1):
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if not any(header):
            empty_header += 1
            fail("第 %d 张表表头为空" % index)
        if len(table.rows) < 2:
            fail("第 %d 张表只有表头没有数据行" % index)
    if not empty_header:
        print("  PASS  %d 张表表头与数据行完整" % len(doc.tables))

    header_rows = sum(1 for t in doc.tables if len(t.rows) >= 2)
    print("  表格数据：%s" % ", ".join("%d行%d列" % (len(t.rows), len(t.columns)) for t in doc.tables[:6]) + " …")

    print("\n结果：%s" % ("全部通过" if ok else "存在问题"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
